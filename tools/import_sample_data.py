#!/usr/bin/env python3
"""Import the workbook's sample hierarchy into the local SQLite database once."""

from __future__ import annotations

import argparse
import sqlite3
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "inital-requirements/project-requirements/DEVELOPING TEST-AIHR-DEV02.xlsx"
DEFAULT_DATABASE = ROOT / "src/backend/Piping.Api/piping.db"
SCALE = Decimal("0.000001")


def normalize_header(value: Any) -> str:
    return " ".join(str(value or "").replace("\n", " ").strip().lower().split())


def read_records(sheet: Any, required: set[str]) -> list[dict[str, Any]]:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        headers = [normalize_header(value) for value in row]
        if required.issubset(set(headers)):
            records = []
            for values in sheet.iter_rows(min_row=row_number + 1, values_only=True):
                record = {header: value for header, value in zip(headers, values) if header and value is not None}
                if any(value is not None for value in values):
                    records.append(record)
            return records
    raise ValueError(f"Could not find header row in worksheet {sheet.title!r}")


def required_value(row: dict[str, Any], key: str, sheet: str, row_number: int) -> Any:
    value = row.get(key)
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"{sheet} row {row_number}: missing {key}")
    return value


def decimal_text(value: Any, sheet: str, row_number: int, field: str) -> str:
    try:
        number = Decimal(str(value))
        rounded = number.quantize(SCALE, rounding=ROUND_HALF_EVEN)
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"{sheet} row {row_number}: invalid decimal {field}={value!r}") from error
    if abs(number - rounded) > Decimal("0.000000001"):
        raise ValueError(f"{sheet} row {row_number}: {field} exceeds the supported 6 decimal places")
    return format(rounded, "f").rstrip("0").rstrip(".") or "0"


def date_text(value: Any, sheet: str, row_number: int, field: str) -> str:
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(str(value)).isoformat()
    except ValueError as error:
        raise ValueError(f"{sheet} row {row_number}: invalid date {field}={value!r}") from error


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE)
    parser.add_argument("--dry-run", action="store_true", help="Validate workbook and relationships without writing")
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    info = read_records(workbook["Sample Info Data"], {"line_number", "pipe_size"})
    cmls = read_records(workbook["Sample CML Data"], {"line_number", "cml_number", "cml_description"})
    points = read_records(workbook["Sample TP Data"], {"line_number", "cml_number", "tp_number", "tp_description"})
    readings = read_records(workbook["Sample Thickness Data "], {"line_number", "cml_number", "tp_number", "inspection_date", "actual_thickness"})

    lines_by_number: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(info, start=1):
        number = str(required_value(row, "line_number", "Sample Info Data", index)).strip()
        if number in lines_by_number:
            raise ValueError(f"Sample Info Data has duplicate line_number {number!r}")
        lines_by_number[number] = row

    cml_by_key: dict[tuple[str, int], dict[str, Any]] = {}
    for index, row in enumerate(cmls, start=1):
        line = str(required_value(row, "line_number", "Sample CML Data", index)).strip()
        key = (line, int(required_value(row, "cml_number", "Sample CML Data", index)))
        if line not in lines_by_number:
            raise ValueError(f"Sample CML Data row {index}: missing parent line {line!r}")
        if key in cml_by_key:
            raise ValueError(f"Sample CML Data has duplicate key {key!r}")
        cml_by_key[key] = row

    point_by_key: dict[tuple[str, int, int], dict[str, Any]] = {}
    for index, row in enumerate(points, start=1):
        line = str(required_value(row, "line_number", "Sample TP Data", index)).strip()
        cml_number = int(required_value(row, "cml_number", "Sample TP Data", index))
        key = (line, cml_number, int(required_value(row, "tp_number", "Sample TP Data", index)))
        if (line, cml_number) not in cml_by_key:
            raise ValueError(f"Sample TP Data row {index}: missing parent CML {(line, cml_number)!r}")
        if key in point_by_key:
            raise ValueError(f"Sample TP Data has duplicate key {key!r}")
        angle = int(required_value(row, "tp_description", "Sample TP Data", index))
        if angle < 0 or angle > 360:
            raise ValueError(f"Sample TP Data row {index}: invalid tp_description {angle}")
        point_by_key[key] = row

    reading_keys: set[tuple[str, int, int, str]] = set()
    for index, row in enumerate(readings, start=1):
        line = str(required_value(row, "line_number", "Sample Thickness Data", index)).strip()
        cml_number = int(required_value(row, "cml_number", "Sample Thickness Data", index))
        tp_number = int(required_value(row, "tp_number", "Sample Thickness Data", index))
        if (line, cml_number, tp_number) not in point_by_key:
            raise ValueError(f"Sample Thickness Data row {index}: missing parent Test Point")
        inspected = date_text(required_value(row, "inspection_date", "Sample Thickness Data", index), "Sample Thickness Data", index, "inspection_date")
        decimal_text(required_value(row, "actual_thickness", "Sample Thickness Data", index), "Sample Thickness Data", index, "actual_thickness")
        key = (line, cml_number, tp_number, inspected)
        if key in reading_keys:
            raise ValueError(f"Sample Thickness Data has duplicate key {key!r}")
        reading_keys.add(key)

    print(f"Validated workbook relationships: {len(info)} lines, {len(cmls)} CMLs, {len(points)} test points, {len(readings)} readings")
    if args.dry_run:
        print("Dry run complete; database was not changed.")
        return

    if not args.database.exists():
        raise FileNotFoundError(f"Database not found: {args.database}. Start the API once to apply migrations first.")
    connection = sqlite3.connect(args.database, timeout=30)
    connection.execute("PRAGMA foreign_keys = ON")
    try:
        with connection:
            table_names = ("PipingLines", "Cmls", "TestPoints", "ThicknessReadings")
            counts = {table: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] for table in table_names}
            if any(counts.values()):
                raise RuntimeError(f"Refusing to import into a non-empty business database: {counts}. Import is all-or-nothing and only supports an empty database.")

            line_ids: dict[str, int] = {}
            for index, (number, row) in enumerate(lines_by_number.items(), start=1):
                values = (
                    number,
                    row.get("location"), row.get("from"), row.get("to"), row.get("drawing number"),
                    row.get("service"), row.get("material"),
                    date_text(row["inservice_date"], "Sample Info Data", index, "inservice_date") if row.get("inservice_date") is not None else None,
                    decimal_text(row["pipe_size"], "Sample Info Data", index, "pipe_size"),
                    decimal_text(row["original_thickness"], "Sample Info Data", index, "original_thickness") if row.get("original_thickness") is not None else None,
                    decimal_text(row["stress"], "Sample Info Data", index, "stress") if row.get("stress") is not None else None,
                    decimal_text(row["joint_efficiency"], "Sample Info Data", index, "joint_efficiency") if row.get("joint_efficiency") is not None else None,
                    decimal_text(row["ca"], "Sample Info Data", index, "ca") if row.get("ca") is not None else None,
                    decimal_text(row["design_life"], "Sample Info Data", index, "design_life") if row.get("design_life") is not None else None,
                    decimal_text(row["design_pressure"], "Sample Info Data", index, "design_pressure") if row.get("design_pressure") is not None else None,
                    decimal_text(row["operating_pressure"], "Sample Info Data", index, "operating_pressure") if row.get("operating_pressure") is not None else None,
                    decimal_text(row["design_temperature"], "Sample Info Data", index, "design_temperature") if row.get("design_temperature") is not None else None,
                    decimal_text(row["operating_temperature"], "Sample Info Data", index, "operating_temperature") if row.get("operating_temperature") is not None else None,
                )
                cursor = connection.execute("""INSERT INTO PipingLines
                    (LineNumber, Location, FromDescription, ToDescription, DrawingNumber, Service, Material, InserviceDate,
                     PipeSizeInch, OriginalThicknessMm, Stress, JointEfficiency, CorrosionAllowanceMm, DesignLife,
                     DesignPressure, OperatingPressure, DesignTemperature, OperatingTemperature)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", values)
                line_ids[number] = cursor.lastrowid

            cml_ids: dict[tuple[str, int], int] = {}
            for (line, number), row in cml_by_key.items():
                cursor = connection.execute("INSERT INTO Cmls (PipingLineId, CmlNumber, Description) VALUES (?, ?, ?)",
                    (line_ids[line], number, str(row["cml_description"]).strip()))
                cml_ids[(line, number)] = cursor.lastrowid

            point_ids: dict[tuple[str, int, int], int] = {}
            for (line, cml_number, number), row in point_by_key.items():
                cursor = connection.execute("INSERT INTO TestPoints (CmlId, TpNumber, TpDescription, Note) VALUES (?, ?, ?, ?)",
                    (cml_ids[(line, cml_number)], number, int(row["tp_description"]), row.get("note")))
                point_ids[(line, cml_number, number)] = cursor.lastrowid

            for index, row in enumerate(readings, start=1):
                key = (str(row["line_number"]).strip(), int(row["cml_number"]), int(row["tp_number"]))
                connection.execute("INSERT INTO ThicknessReadings (TestPointId, InspectionDate, ActualThicknessMm) VALUES (?, ?, ?)",
                    (point_ids[key], date_text(row["inspection_date"], "Sample Thickness Data", index, "inspection_date"), decimal_text(row["actual_thickness"], "Sample Thickness Data", index, "actual_thickness")))

        print("Import committed successfully:")
        for table in table_names:
            count = connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
            print(f"  {table}: {count}")
    finally:
        connection.close()


if __name__ == "__main__":
    main()
